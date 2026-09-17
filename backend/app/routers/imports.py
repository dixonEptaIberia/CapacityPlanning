"""Master-data import endpoints (R11).

Imports cadence master data from an Excel file so unstable technical values can
be updated without changing application code. Expected columns (first sheet):
``label``, ``workers``, ``shifts``, ``pieces_per_week``.
"""
from __future__ import annotations

import io

from fastapi import APIRouter, Depends, UploadFile
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.auth import require_role
from app.db.base import get_db
from app.models import CadenceOption, RoleName, User
from app.schemas import ImportResult

router = APIRouter(tags=["imports"])

_REQUIRED_COLUMNS = ["label", "workers", "shifts", "pieces_per_week"]


@router.post("/imports/master-data", response_model=ImportResult)
async def import_master_data(
    file: UploadFile,
    db: Session = Depends(get_db),
    _: User = Depends(require_role(RoleName.central_admin)),
):
    """Import/update cadence options from Excel; reject invalid files (R11.1-R11.3)."""
    content = await file.read()
    errors: list[str] = []
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception:  # noqa: BLE001 - surface a clean validation error
        return ImportResult(imported=0, updated=0, errors=["File is not a valid Excel workbook."])

    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return ImportResult(imported=0, updated=0, errors=["The spreadsheet is empty."])

    header = [str(c).strip().lower() if c is not None else "" for c in rows[0]]
    missing = [c for c in _REQUIRED_COLUMNS if c not in header]
    if missing:
        return ImportResult(
            imported=0,
            updated=0,
            errors=[f"Missing required column(s): {', '.join(missing)}."],
        )
    idx = {c: header.index(c) for c in _REQUIRED_COLUMNS}

    imported = 0
    updated = 0
    # Validate all rows first; do not corrupt existing data on failure (R11.3).
    parsed: list[dict] = []
    for r, row in enumerate(rows[1:], start=2):
        try:
            label = str(row[idx["label"]]).strip()
            workers = int(row[idx["workers"]])
            shifts = int(row[idx["shifts"]])
            pieces = int(row[idx["pieces_per_week"]])
        except (TypeError, ValueError):
            errors.append(f"Row {r}: invalid or missing numeric values.")
            continue
        if not label:
            errors.append(f"Row {r}: label is required.")
            continue
        parsed.append(
            {"label": label, "workers": workers, "shifts": shifts, "pieces_per_week": pieces}
        )

    if errors:
        return ImportResult(imported=0, updated=0, errors=errors)

    for item in parsed:
        existing = db.scalar(select(CadenceOption).where(CadenceOption.label == item["label"]))
        if existing:
            existing.workers = item["workers"]
            existing.shifts = item["shifts"]
            existing.pieces_per_week = item["pieces_per_week"]
            existing.active = True
            updated += 1
        else:
            db.add(CadenceOption(active=True, **item))
            imported += 1

    db.commit()
    return ImportResult(imported=imported, updated=updated, errors=[])
